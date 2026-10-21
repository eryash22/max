import io

import openpyxl

from app.models.contracts import AmcContract, AmcContractAsset, CmcContract, Warranty
from app.models.device import Asset
from app.models.maintenance import CalibrationRecord, CalibrationSchedule, PMRecord, PMSchedule
from app.models.vendor import Vendor, VendorContact


def _build_workbook(*, vendor_id_for_amc="V-DEMO-01", vendor_id_for_contact="V-DEMO-01") -> bytes:
    wb = openpyxl.Workbook()

    hospital = wb.active
    hospital.title = "Hospital Master"
    hospital.append(["Hospital master"])
    hospital.append(["Hospital Code*", "Hospital Name*", "Status*"])
    hospital.append(["MAX-DEMO-01", "DEMO Hospital", "ACTIVE"])

    device_models = wb.create_sheet("Device Model Master")
    device_models.append(["Catalog"])
    device_models.append(["Manufacturer*", "Model Number*", "Device Name*", "Category*"])
    device_models.append(["DEMO Manufacturer", "MODEL-001", "DEMO Patient Monitor", "PATIENT_MONITOR"])

    devices = wb.create_sheet("Device Master")
    devices.append(["Devices"])
    devices.append([
        "Asset Number*", "Device Name*", "Category*", "Manufacturer*", "Model Number*",
        "Hospital Code*", "Status*",
    ])
    devices.append([
        "ASSET-DEMO-0001", "DEMO Patient Monitor", "PATIENT_MONITOR", "DEMO Manufacturer",
        "MODEL-001", "MAX-DEMO-01", "OPERATIONAL",
    ])

    vendor_master = wb.create_sheet("Vendor Master")
    vendor_master.append(["Vendors"])
    vendor_master.append(["Vendor ID*", "Vendor Name*", "Vendor Category*", "Status*", "Email", "Phone", "City", "SLA Response Hours"])
    vendor_master.append(["V-DEMO-01", "DEMO Vendor Co", "OEM", "ACTIVE", "vendor@example.invalid", "+911234567890", "New Delhi", 24])

    vendor_contacts = wb.create_sheet("Vendor Contacts")
    vendor_contacts.append(["Contacts"])
    vendor_contacts.append(["Vendor Contact ID*", "Vendor ID*", "Contact Name*", "Email", "Phone", "Role", "Status*"])
    vendor_contacts.append(["VC-DEMO-001", vendor_id_for_contact, "Demo Contact", "contact@example.invalid", "+911234567891", "SERVICE", "ACTIVE"])

    pm_schedule = wb.create_sheet("PM Schedule")
    pm_schedule.append(["PM Schedules"])
    pm_schedule.append(["PM Schedule ID*", "Asset Number*", "Frequency*", "Next Due Date*", "Status*"])
    pm_schedule.append(["PMS-DEMO-001", "ASSET-DEMO-0001", "QUARTERLY", "2026-12-01", "ACTIVE"])

    pm_records = wb.create_sheet("PM Records")
    pm_records.append(["PM Records"])
    pm_records.append(["PM Record ID*", "Asset Number*", "Performed Date*", "Result*", "Performed By*", "Next Due Date", "Notes"])
    pm_records.append(["PMR-DEMO-001", "ASSET-DEMO-0001", "2026-01-01", "PASS", "demo.engineer@example.invalid", "2026-04-01", "DEMO record"])

    calibration_schedule = wb.create_sheet("Calibration Schedule")
    calibration_schedule.append(["Calibration Schedules"])
    calibration_schedule.append(["Calibration Schedule ID*", "Asset Number*", "Frequency*", "Next Due Date*", "Status*"])
    calibration_schedule.append(["CS-DEMO-001", "ASSET-DEMO-0001", "ANNUAL", "2027-01-01", "ACTIVE"])

    calibration_records = wb.create_sheet("Calibration Records")
    calibration_records.append(["Calibration Records"])
    calibration_records.append(["Calibration Record ID*", "Asset Number*", "Performed Date*", "Result*", "Certificate Number", "Provider/Vendor ID", "Next Due Date", "Notes"])
    calibration_records.append(["CR-DEMO-001", "ASSET-DEMO-0001", "2026-01-01", "PASS", "CERT-DEMO-001", "V-DEMO-01", "2027-01-01", "DEMO record"])

    amc = wb.create_sheet("AMC Contracts")
    amc.append(["AMC"])
    amc.append(["Contract Number*", "Asset Number*", "Vendor ID*", "Start Date*", "End Date*", "Status*", "SLA Response Hours", "Cost"])
    amc.append(["AMC-DEMO-001", "ASSET-DEMO-0001", vendor_id_for_amc, "2026-01-01", "2026-12-31", "ACTIVE", 24, 5000])

    cmc = wb.create_sheet("CMC Contracts")
    cmc.append(["CMC"])
    cmc.append(["Contract Number*", "Asset Number*", "Vendor ID*", "Start Date*", "End Date*", "Status*", "SLA Response Hours", "Cost"])
    cmc.append(["CMC-DEMO-001", "ASSET-DEMO-0001", vendor_id_for_amc, "2026-01-01", "2026-12-31", "ACTIVE", 24, 8000])

    warranty = wb.create_sheet("Warranty")
    warranty.append(["Warranty"])
    warranty.append(["Warranty ID*", "Asset Number*", "Vendor ID", "Start Date*", "End Date*", "Status*", "Terms"])
    warranty.append(["W-DEMO-01", "ASSET-DEMO-0001", vendor_id_for_amc, "2025-01-01", "2026-12-31", "ACTIVE", "Standard terms"])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload(admin_client, **kwargs):
    file_bytes = _build_workbook(**kwargs)
    response = admin_client.post(
        "/api/v1/imports",
        files={"file": ("test_import_module14.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _approve(admin_client, job):
    approval = admin_client.post(f"/api/v1/imports/{job['id']}/approve", json={"decision": "APPROVED"})
    assert approval.status_code == 200, approval.text
    return approval.json()


def test_upload_recognizes_all_module14_sheets(admin_client):
    job = _upload(admin_client)
    assert job["summary"]["unsupported_sheets"] == []
    mapping = job["column_mapping"]
    assert mapping["Vendor Master"] == "vendor"
    assert mapping["Vendor Contacts"] == "vendor_contact"
    assert mapping["PM Schedule"] == "pm_schedule"
    assert mapping["PM Records"] == "pm_record"
    assert mapping["Calibration Schedule"] == "calibration_schedule"
    assert mapping["Calibration Records"] == "calibration_record"
    assert mapping["AMC Contracts"] == "amc_contract"
    assert mapping["CMC Contracts"] == "cmc_contract"
    assert mapping["Warranty"] == "warranty"


def test_approve_creates_vendor(admin_client, db_session):
    job = _upload(admin_client)
    _approve(admin_client, job)
    vendor = db_session.query(Vendor).filter(Vendor.name == "DEMO Vendor Co").first()
    assert vendor is not None
    assert vendor.vendor_type == "OEM"
    assert vendor.status == "ACTIVE"
    assert vendor.response_sla_hours == 24


def test_approve_creates_vendor_contact_resolved_to_real_vendor(admin_client, db_session):
    job = _upload(admin_client)
    _approve(admin_client, job)
    vendor = db_session.query(Vendor).filter(Vendor.name == "DEMO Vendor Co").first()
    contact = db_session.query(VendorContact).filter(VendorContact.name == "Demo Contact").first()
    assert contact is not None
    assert contact.vendor_id == vendor.id


def test_approve_creates_pm_schedule_and_record_linked_to_asset(admin_client, db_session):
    job = _upload(admin_client)
    _approve(admin_client, job)
    asset = db_session.query(Asset).filter(Asset.asset_number == "ASSET-DEMO-0001").first()
    assert asset is not None

    schedule = db_session.query(PMSchedule).filter(PMSchedule.asset_id == asset.id).first()
    assert schedule is not None
    assert schedule.frequency == "QUARTERLY"
    assert schedule.is_active is True

    record = db_session.query(PMRecord).filter(PMRecord.asset_id == asset.id).first()
    assert record is not None
    assert record.performed_date == "2026-01-01"
    assert record.functional_check_result == "PASS"
    assert record.schedule_id == schedule.id


def test_approve_normalizes_annual_frequency_to_yearly(admin_client, db_session):
    # Calibration Schedule sheet uses "ANNUAL" (the Phase 0 template's Reference
    # Values wording); the actual Module 03A vocabulary is YEARLY.
    job = _upload(admin_client)
    _approve(admin_client, job)
    asset = db_session.query(Asset).filter(Asset.asset_number == "ASSET-DEMO-0001").first()
    schedule = db_session.query(CalibrationSchedule).filter(CalibrationSchedule.asset_id == asset.id).first()
    assert schedule is not None
    assert schedule.frequency == "YEARLY"


def test_approve_creates_calibration_record_with_pass_fail_derived(admin_client, db_session):
    job = _upload(admin_client)
    _approve(admin_client, job)
    asset = db_session.query(Asset).filter(Asset.asset_number == "ASSET-DEMO-0001").first()
    record = db_session.query(CalibrationRecord).filter(CalibrationRecord.asset_id == asset.id).first()
    assert record is not None
    assert record.result == "PASS"
    assert record.pass_fail == "PASS"
    assert record.calibration_agency == "V-DEMO-01"  # free text, not resolved (no vendor_id column)
    assert record.certificate_number == "CERT-DEMO-001"


def test_approve_creates_amc_and_cmc_contracts_with_resolved_vendor_and_asset_link(admin_client, db_session):
    job = _upload(admin_client)
    _approve(admin_client, job)
    vendor = db_session.query(Vendor).filter(Vendor.name == "DEMO Vendor Co").first()
    asset = db_session.query(Asset).filter(Asset.asset_number == "ASSET-DEMO-0001").first()

    amc = db_session.query(AmcContract).filter(AmcContract.contract_number == "AMC-DEMO-001").first()
    assert amc is not None
    assert amc.vendor_id == vendor.id
    assert amc.contract_value == 5000
    link = db_session.query(AmcContractAsset).filter(
        AmcContractAsset.amc_contract_id == amc.id, AmcContractAsset.asset_id == asset.id
    ).first()
    assert link is not None

    cmc = db_session.query(CmcContract).filter(CmcContract.contract_number == "CMC-DEMO-001").first()
    assert cmc is not None
    assert cmc.vendor_id == vendor.id


def test_approve_creates_warranty_one_per_asset(admin_client, db_session):
    job = _upload(admin_client)
    _approve(admin_client, job)
    vendor = db_session.query(Vendor).filter(Vendor.name == "DEMO Vendor Co").first()
    asset = db_session.query(Asset).filter(Asset.asset_number == "ASSET-DEMO-0001").first()

    warranty = db_session.query(Warranty).filter(Warranty.asset_id == asset.id).first()
    assert warranty is not None
    assert warranty.vendor_id == vendor.id
    assert warranty.warranty_number == "W-DEMO-01"
    assert warranty.coverage_notes == "Standard terms"


def test_unresolved_vendor_on_contract_warns_but_does_not_block(admin_client, db_session):
    # A "Vendor ID" that matches neither this job's own vendor master sheet nor any
    # existing vendor's name should not fail the whole row — vendor_id is nullable on
    # AmcContract/CmcContract/Warranty (unlike VendorContact, where it's required).
    job = _upload(admin_client, vendor_id_for_amc="UNKNOWN-VENDOR-999")
    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 0

    amc = db_session.query(AmcContract).filter(AmcContract.contract_number == "AMC-DEMO-001").first()
    assert amc is not None
    assert amc.vendor_id is None

    preview = admin_client.get(f"/api/v1/imports/{job['id']}/preview").json()
    assert any(
        e["severity"] == "WARNING" and "vendor" in (e["field"] or "").lower() for e in preview["errors"]
    )


def test_unresolved_asset_on_pm_schedule_fails_that_row_only(admin_client):
    wb = openpyxl.Workbook()
    hospital = wb.active
    hospital.title = "Hospital Master"
    hospital.append(["Hospital master"])
    hospital.append(["Hospital Code*", "Hospital Name*", "Status*"])
    hospital.append(["MAX-DEMO-02", "DEMO Hospital 2", "ACTIVE"])

    pm_schedule = wb.create_sheet("PM Schedule")
    pm_schedule.append(["PM Schedules"])
    pm_schedule.append(["PM Schedule ID*", "Asset Number*", "Frequency*", "Next Due Date*", "Status*"])
    pm_schedule.append(["PMS-DEMO-002", "DOES-NOT-EXIST", "MONTHLY", "2026-12-01", "ACTIVE"])

    buffer = io.BytesIO()
    wb.save(buffer)
    response = admin_client.post(
        "/api/v1/imports",
        files={"file": ("test_bad_asset.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    job = response.json()
    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 1

    preview = admin_client.get(f"/api/v1/imports/{job['id']}/preview").json()
    assert any(e["severity"] == "ERROR" and e["field"] == "asset_number" for e in preview["errors"])


def test_reimporting_same_vendor_updates_not_duplicates(admin_client, db_session):
    job1 = _upload(admin_client)
    _approve(admin_client, job1)
    job2 = _upload(admin_client)
    result2 = _approve(admin_client, job2)

    vendors = db_session.query(Vendor).filter(Vendor.name == "DEMO Vendor Co").all()
    assert len(vendors) == 1
    assert result2["summary"]["updated"] >= 1
