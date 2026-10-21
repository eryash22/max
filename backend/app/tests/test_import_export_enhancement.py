"""Tests for the Import & Export Enhancement task: CSV import support, the six newly
wired datasets (Departments, Locations, Inventory Locations, Spare Parts/Inventory
Items, Service History, Document Metadata), and the reusable export system.

Import flow (upload -> validate -> preview -> approval -> import -> audit) is verified
unchanged — these tests exercise it against new inputs, they don't change its shape.
"""

import csv
import io

import openpyxl

from app.models.audit import AuditLog
from app.models.device import Asset
from app.models.documents import Document
from app.models.hospital import Branch, Department, Location
from app.models.incident import ServiceVisit
from app.models.inventory import Consumable, InventoryItem, InventoryLocation
from app.models.review import ReviewItem

from .conftest import auth_headers, create_user


def _csv_bytes(rows: list[list[object]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def _upload_csv(admin_client, filename: str, rows: list[list[object]]):
    response = admin_client.post(
        "/api/v1/imports",
        files={"file": (filename, _csv_bytes(rows), "text/csv")},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _approve(admin_client, job):
    response = admin_client.post(f"/api/v1/imports/{job['id']}/approve", json={"decision": "APPROVED"})
    assert response.status_code == 200, response.text
    return response.json()


def _seed_branch_and_asset(admin_client) -> tuple[dict, dict]:
    """Build one branch + device model + asset via the existing .xlsx pipeline (already
    covered by test_imports.py/test_imports_module14.py) so CSV-only tests can focus on
    the CSV-specific behavior instead of re-deriving base data."""
    wb = openpyxl.Workbook()
    hospital = wb.active
    hospital.title = "Hospital Master"
    hospital.append(["Hospital master"])
    hospital.append(["Hospital Code*", "Hospital Name*", "Status*"])
    hospital.append(["CSV-DEMO-01", "CSV Demo Hospital", "ACTIVE"])

    device_models = wb.create_sheet("Device Model Master")
    device_models.append(["Catalog"])
    device_models.append(["Manufacturer*", "Model Number*", "Device Name*", "Category*"])
    device_models.append(["CSV Manufacturer", "CSV-MODEL-1", "CSV Patient Monitor", "PATIENT_MONITOR"])

    devices = wb.create_sheet("Device Master")
    devices.append(["Devices"])
    devices.append(["Asset Number*", "Device Name*", "Category*", "Manufacturer*", "Model Number*", "Hospital Code*", "Status*"])
    devices.append(["CSV-ASSET-0001", "CSV Patient Monitor", "PATIENT_MONITOR", "CSV Manufacturer", "CSV-MODEL-1", "CSV-DEMO-01", "OPERATIONAL"])

    buffer = io.BytesIO()
    wb.save(buffer)
    response = admin_client.post(
        "/api/v1/imports",
        files={"file": ("seed.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    job = response.json()
    _approve(admin_client, job)
    return {"code": "CSV-DEMO-01"}, {"asset_number": "CSV-ASSET-0001"}


# ---- CSV import: file-type detection, headers, mapping, validation ----

def test_import_csv_by_filename_match(admin_client, db_session):
    job = _upload_csv(
        admin_client, "hospital_master.csv",
        [["Hospital Code*", "Hospital Name*", "Status*"], ["CSVFN-01", "CSV Filename Hospital", "ACTIVE"]],
    )
    assert job["status"] == "READY_FOR_REVIEW"
    assert job["column_mapping"]["hospital_master"] == "branch"

    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 0
    branch = db_session.query(Branch).filter(Branch.code == "CSVFN-01").first()
    assert branch is not None
    assert branch.name == "CSV Filename Hospital"


def test_import_csv_best_fit_detection_when_filename_unrecognized(admin_client, db_session):
    # A generically-named file ("export.csv") can't be matched by filename, so the
    # importer falls back to scoring header overlap against known sheet configs —
    # these headers uniquely identify "vendor master".
    job = _upload_csv(
        admin_client, "export.csv",
        [
            ["Vendor ID*", "Vendor Name*", "Vendor Category*", "Status*", "Email", "Phone", "City", "SLA Response Hours"],
            ["V-CSV-01", "CSV Best-Fit Vendor", "OEM", "ACTIVE", "v@example.invalid", "+911111111111", "Pune", 12],
        ],
    )
    assert job["column_mapping"]["export"] == "vendor"

    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 0


def test_import_csv_unrecognized_headers_stage_only_no_crash(admin_client):
    job = _upload_csv(admin_client, "mystery.csv", [["Foo", "Bar"], ["1", "2"]])
    assert job["status"] == "READY_FOR_REVIEW"
    assert job["summary"]["unsupported_sheets"] == ["mystery"]
    assert job["column_mapping"]["mystery"] is None


def test_xlsx_still_accepted_after_csv_support_added(admin_client):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Hospital Master"
    sheet.append(["Hospital master"])
    sheet.append(["Hospital Code*", "Hospital Name*", "Status*"])
    sheet.append(["XLSX-01", "Still Works Hospital", "ACTIVE"])
    buffer = io.BytesIO()
    wb.save(buffer)
    response = admin_client.post(
        "/api/v1/imports",
        files={"file": ("still_works.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201
    assert response.json()["status"] == "READY_FOR_REVIEW"


def test_import_rejects_unsupported_extension(admin_client):
    response = admin_client.post(
        "/api/v1/imports",
        files={"file": ("data.txt", b"not a spreadsheet", "text/plain")},
    )
    assert response.status_code == 400


def test_import_rejects_oversized_file(admin_client):
    from app.routers.imports import MAX_IMPORT_UPLOAD_BYTES

    oversized = b"a" * (MAX_IMPORT_UPLOAD_BYTES + 1)
    response = admin_client.post(
        "/api/v1/imports",
        files={"file": ("big.csv", oversized, "text/csv")},
    )
    assert response.status_code == 413


def test_import_filename_is_sanitized_on_disk(admin_client):
    """Security hardening pass — the on-disk filename previously used the raw,
    unsanitized upload filename (unlike documents.py's create_version, which already
    sanitizes). A path-traversal-shaped filename must not escape storage/imports."""
    from app.config import settings

    response = admin_client.post(
        "/api/v1/imports",
        files={"file": ("../../etc/passwd.csv", b"Name\r\ntest\r\n", "text/csv")},
    )
    assert response.status_code == 201, response.text
    job_id = response.json()["id"]

    imports_dir = settings.file_storage_path / "imports"
    matches = [p for p in imports_dir.iterdir() if p.name.startswith(job_id)]
    assert len(matches) == 1
    saved_path = matches[0]
    # Resolved path must still be inside storage/imports — not escaped via ".." — and
    # the stored filename itself must carry no path separators.
    assert saved_path.resolve().parent == imports_dir.resolve()
    assert "/" not in saved_path.name and ".." not in saved_path.name


# ---- New datasets: Departments, Locations, Inventory Locations, Spare Parts,
# Service History, Document Metadata ----

def test_import_department_master(admin_client, db_session):
    branch, _ = _seed_branch_and_asset(admin_client)
    job = _upload_csv(
        admin_client, "department_master.csv",
        [["Department Code*", "Department Name*", "Hospital Code*", "Status*"], ["ICU-1", "ICU", branch["code"], "ACTIVE"]],
    )
    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 0
    dept = db_session.query(Department).filter(Department.code == "ICU-1").first()
    assert dept is not None
    assert dept.name == "ICU"


def test_import_location_master_requires_department_first(admin_client, db_session):
    branch, _ = _seed_branch_and_asset(admin_client)
    dept_job = _upload_csv(
        admin_client, "department_master.csv",
        [["Department Code*", "Department Name*", "Hospital Code*", "Status*"], ["ICU-2", "ICU Two", branch["code"], "ACTIVE"]],
    )
    _approve(admin_client, dept_job)

    loc_job = _upload_csv(
        admin_client, "location_master.csv",
        [
            ["Location Code*", "Location Name*", "Hospital Code*", "Department Code*", "Room/Bed", "Status*"],
            ["ICU2-BED-01", "ICU2 Bed 01", branch["code"], "ICU-2", "Bed 01", "ACTIVE"],
        ],
    )
    result = _approve(admin_client, loc_job)
    assert result["summary"]["failed"] == 0
    location = db_session.query(Location).filter(Location.code == "ICU2-BED-01").first()
    assert location is not None
    assert location.room == "Bed 01"


def test_import_location_master_unresolved_department_fails_row(admin_client, db_session):
    branch, _ = _seed_branch_and_asset(admin_client)
    job = _upload_csv(
        admin_client, "location_master.csv",
        [
            ["Location Code*", "Location Name*", "Hospital Code*", "Department Code*", "Status*"],
            ["ORPHAN-BED", "Orphan Bed", branch["code"], "DOES-NOT-EXIST", "ACTIVE"],
        ],
    )
    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 1
    preview = admin_client.get(f"/api/v1/imports/{job['id']}/preview").json()
    assert any(e["severity"] == "ERROR" and e["field"] == "department_code" for e in preview["errors"])


def test_import_inventory_locations(admin_client, db_session):
    branch, _ = _seed_branch_and_asset(admin_client)
    job = _upload_csv(
        admin_client, "inventory_locations.csv",
        [["Location Code*", "Location Name*", "Hospital Code*", "Kind"], ["WH-CSV-1", "CSV Warehouse", branch["code"], "WAREHOUSE"]],
    )
    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 0
    loc = db_session.query(InventoryLocation).filter(InventoryLocation.code == "WH-CSV-1").first()
    assert loc is not None
    assert loc.kind == "WAREHOUSE"


def test_import_spare_parts(admin_client, db_session):
    job = _upload_csv(
        admin_client, "spare_parts.csv",
        [
            ["Part Number*", "Part Name*", "Category*", "Manufacturer", "Unit*", "Minimum Stock", "Reorder Level"],
            ["PART-CSV-001", "CSV Sensor Cable", "SPARE_PART", "CSV Manufacturer", "EACH", 1, 2],
        ],
    )
    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 0
    item = db_session.query(InventoryItem).filter(InventoryItem.item_code == "PART-CSV-001").first()
    assert item is not None
    assert item.is_consumable is False
    assert item.supplier == "CSV Manufacturer"


def test_import_service_history(admin_client, db_session):
    _, asset = _seed_branch_and_asset(admin_client)
    job = _upload_csv(
        admin_client, "service_history.csv",
        [
            ["Asset Number*", "Event Date*", "Service Type*", "Complaint", "Root Cause", "Resolution Status*", "Cost"],
            [asset["asset_number"], "2026-01-15", "CORRECTIVE", "Screen flickering", "Loose cable", "RESOLVED", 500],
        ],
    )
    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 0
    visit = db_session.query(ServiceVisit).filter(ServiceVisit.visit_date == "2026-01-15").first()
    assert visit is not None
    assert "Screen flickering" in visit.summary
    assert "Loose cable" in visit.summary
    assert visit.cost == 500


def test_import_service_history_unresolved_asset_fails_row(admin_client):
    job = _upload_csv(
        admin_client, "service_history.csv",
        [["Asset Number*", "Event Date*", "Service Type*", "Resolution Status*"], ["NOT-REAL-ASSET", "2026-01-01", "PM", "RESOLVED"]],
    )
    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 1


def test_import_document_metadata(admin_client, db_session):
    job = _upload_csv(
        admin_client, "document_metadata.csv",
        [
            ["Document Title*", "Document Type*", "Manufacturer", "Model Number", "Access Level"],
            ["CSV Service Manual", "SERVICE_MANUAL", "CSV Manufacturer", "CSV-MODEL-1", "INTERNAL"],
        ],
    )
    result = _approve(admin_client, job)
    assert result["summary"]["failed"] == 0
    doc = db_session.query(Document).filter(Document.title == "CSV Service Manual").first()
    assert doc is not None
    assert doc.document_type == "SERVICE_MANUAL"


# ---- Duplicate detection generalizes to new sheet types without new code ----

def test_duplicate_detection_on_new_sheet_type(admin_client):
    branch, _ = _seed_branch_and_asset(admin_client)
    job = _upload_csv(
        admin_client, "department_master.csv",
        [
            ["Department Code*", "Department Name*", "Hospital Code*", "Status*"],
            ["DUP-1", "Dup Dept", branch["code"], "ACTIVE"],
            ["DUP-1", "Dup Dept", branch["code"], "ACTIVE"],
        ],
    )
    assert job["summary"]["duplicate_candidates"] >= 1


# ---- Review Queue behavior is unchanged for the entity type that already used it ----

def test_review_queue_still_populated_for_misplaced_assignments(admin_client, db_session):
    wb = openpyxl.Workbook()
    hospital = wb.active
    hospital.title = "Hospital Master"
    hospital.append(["Hospital master"])
    hospital.append(["Hospital Code*", "Hospital Name*", "Status*"])
    hospital.append(["RQ-01", "Review Queue Hospital", "ACTIVE"])

    assignments = wb.create_sheet("Branch Device Assignments")
    assignments.append(["Assignments"])
    assignments.append([
        "Assignment ID", "Branch Name*", "Catalog Record Name", "Manufacturer", "Model Number",
        "Known Quantity", "Assignment Status*", "Inventory Mapping Status", "Asset Identification Status",
        "Source Name*", "Data Quality Flag", "Notes", "Environment",
    ])
    assignments.append([
        None, "Review Queue Hospital", "Misplaced Device", None, None, None,
        "LOCATION_ASSET_VERIFICATION_REQUIRED", "Data Received", "PENDING_ASSET_IDENTIFICATION",
        "Misplaced Device", "MISPLACED_REPORTED_BY_BIOMED", "Reported misplaced", "DEVELOPMENT / DEMO DATA",
    ])
    buffer = io.BytesIO()
    wb.save(buffer)
    response = admin_client.post(
        "/api/v1/imports",
        files={"file": ("rq.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    job = response.json()
    _approve(admin_client, job)

    review_item = db_session.query(ReviewItem).filter(ReviewItem.import_job_id == job["id"]).first()
    assert review_item is not None
    assert review_item.reason_code == "MISPLACED_REPORTED_BY_BIOMED"
    assert review_item.status == "OPEN"


def test_import_audit_log_written_for_csv_approval(admin_client, db_session):
    job = _upload_csv(
        admin_client, "hospital_master.csv",
        [["Hospital Code*", "Hospital Name*", "Status*"], ["AUD-01", "Audit Hospital", "ACTIVE"]],
    )
    _approve(admin_client, job)
    log = db_session.query(AuditLog).filter(AuditLog.import_job_id == job["id"], AuditLog.action == "APPROVE").first()
    assert log is not None


def test_csv_import_requires_permission(client, db_session):
    create_user(db_session, "csvviewer@example.com", "VIEWER", "Password123!")
    headers = auth_headers(client, "csvviewer@example.com", "Password123!")
    response = client.post(
        "/api/v1/imports",
        files={"file": ("x.csv", _csv_bytes([["a"], ["1"]]), "text/csv")},
        headers=headers,
    )
    assert response.status_code == 403


def test_consumables_import_sets_opening_quantity(admin_client, db_session):
    """Regression test for the master-inventory-import opening_quantity gap: the
    Consumables sheet previously had no column wired to Consumable.opening_quantity
    at all (mapping.py's SheetConfig omitted it and the upsert handler never set it).
    Covers both the create path and the update path (re-import must overwrite, not
    just set once)."""
    job = _upload_csv(
        admin_client, "consumables.csv",
        [
            ["Item Name*", "Category*", "Opening Quantity"],
            ["Dysphagia Electrode Pack of 2", "Physiotherapy", "358"],
        ],
    )
    _approve(admin_client, job)
    consumable = db_session.query(Consumable).filter(Consumable.name == "Dysphagia Electrode Pack of 2").first()
    assert consumable is not None
    assert consumable.opening_quantity == 358
    assert consumable.category == "Physiotherapy"

    # Re-import with an updated quantity must overwrite the existing record, not duplicate it.
    job2 = _upload_csv(
        admin_client, "consumables.csv",
        [
            ["Item Name*", "Category*", "Opening Quantity"],
            ["Dysphagia Electrode Pack of 2", "Physiotherapy", "400"],
        ],
    )
    _approve(admin_client, job2)
    db_session.expire_all()
    matches = db_session.query(Consumable).filter(Consumable.name == "Dysphagia Electrode Pack of 2").all()
    assert len(matches) == 1
    assert matches[0].opening_quantity == 400
