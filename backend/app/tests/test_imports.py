import io

import openpyxl

from app.models.device import BranchDeviceAssignment, DeviceModel
from app.models.hospital import Branch
from app.models.review import ReviewItem


def _build_workbook() -> bytes:
    wb = openpyxl.Workbook()

    hospital = wb.active
    hospital.title = "Hospital Master"
    hospital.append(["Hospital master", None, None])
    hospital.append(["Hospital Code", "Hospital Name*", "Status*", "Region"])
    hospital.append([None, "Saket", "ACTIVE", "Delhi/NCR"])
    hospital.append([None, "Noida / Jaypee Max", "ACTIVE", "Delhi/NCR"])

    device_models = wb.create_sheet("Device Model Master")
    device_models.append(["Catalog"])
    device_models.append(["Manufacturer*", "Model Number", "Device Name*", "Category*"])
    device_models.append(["Enraf-Nonius", None, "Enraf Ultrasound", "Ultrasound Therapy"])
    device_models.append(["Chattanooga", "2782", "Chattanooga Intelect Transport Ultrasound", "Ultrasound Therapy"])

    assignments = wb.create_sheet("Branch Device Assignments")
    assignments.append(["Assignments"])
    assignments.append([
        "Assignment ID", "Branch Name*", "Catalog Record Name", "Manufacturer", "Model Number",
        "Known Quantity", "Assignment Status*", "Inventory Mapping Status", "Asset Identification Status",
        "Source Name*", "Data Quality Flag", "Notes", "Environment",
    ])
    # Normal row
    assignments.append([
        None, "Saket", "Chattanooga Intelect Transport Ultrasound", "Chattanooga", "2782", None,
        "PENDING_ASSET_IDENTIFICATION", "Data Received", "PENDING_ASSET_IDENTIFICATION",
        "Chattanooga Intelect Transport Ultrasound", "NONE", None, "DEVELOPMENT / DEMO DATA",
    ])
    # Duplicate of the row above (same branch + source name)
    assignments.append([
        None, "Saket", "Chattanooga Intelect Transport Ultrasound", "Chattanooga", "2782", None,
        "PENDING_ASSET_IDENTIFICATION", "Data Received", "PENDING_ASSET_IDENTIFICATION",
        "Chattanooga Intelect Transport Ultrasound", "NONE", None, "DEVELOPMENT / DEMO DATA",
    ])
    # Noida misplaced device — must be preserved verbatim, not auto-resolved
    assignments.append([
        None, "Noida / Jaypee Max", "Enraf Ultrasound", "Enraf-Nonius", None, None,
        "LOCATION_ASSET_VERIFICATION_REQUIRED", "Data Received", "PENDING_ASSET_IDENTIFICATION",
        "Enraf Ultrasound", "MISPLACED_REPORTED_BY_BIOMED", "Reported misplaced; do not treat as operational or lost.",
        "DEVELOPMENT / DEMO DATA",
    ])
    # Missing required field (no Branch Name)
    assignments.append([
        None, None, None, None, None, None, "REVIEW_QUEUE", "Data Received", "PENDING_ASSET_IDENTIFICATION",
        "Unmapped Source", "UNRESOLVED_SOURCE_NAME", None, "DEVELOPMENT / DEMO DATA",
    ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload(admin_client):
    file_bytes = _build_workbook()
    response = admin_client.post(
        "/api/v1/imports",
        files={"file": ("test_import.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_upload_parses_sheets_and_flags_issues(admin_client):
    job = _upload(admin_client)
    assert job["status"] == "READY_FOR_REVIEW"
    assert job["summary"]["errors"] >= 1  # the missing-branch-name row
    assert job["summary"]["duplicate_candidates"] >= 1  # the duplicate assignment row


def test_preview_lists_rows_and_errors(admin_client):
    job = _upload(admin_client)
    preview = admin_client.get(f"/api/v1/imports/{job['id']}/preview")
    assert preview.status_code == 200
    body = preview.json()
    assert len(body["rows"]) > 0
    assert any(e["severity"] == "ERROR" for e in body["errors"])


def test_approve_creates_entities_and_preserves_misplaced_status(admin_client, db_session):
    job = _upload(admin_client)
    approval = admin_client.post(f"/api/v1/imports/{job['id']}/approve", json={"decision": "APPROVED"})
    assert approval.status_code == 200
    assert approval.json()["status"] == "COMPLETED"

    saket = db_session.query(Branch).filter(Branch.name == "Saket").first()
    assert saket is not None

    device_model = db_session.query(DeviceModel).filter(DeviceModel.name == "Enraf Ultrasound").first()
    assert device_model is not None

    misplaced = db_session.query(BranchDeviceAssignment).filter(
        BranchDeviceAssignment.source_name == "Enraf Ultrasound"
    ).first()
    assert misplaced is not None
    assert misplaced.assignment_status == "LOCATION_ASSET_VERIFICATION_REQUIRED"
    assert misplaced.data_quality_flag == "MISPLACED_REPORTED_BY_BIOMED"

    review_item = db_session.query(ReviewItem).filter(ReviewItem.source_id == misplaced.id).first()
    assert review_item is not None
    assert review_item.reason_code == "MISPLACED_REPORTED_BY_BIOMED"
    assert review_item.status == "OPEN"


def test_approve_fails_when_job_not_in_review(admin_client):
    job = _upload(admin_client)
    admin_client.post(f"/api/v1/imports/{job['id']}/approve", json={"decision": "APPROVED"})
    second_attempt = admin_client.post(f"/api/v1/imports/{job['id']}/approve", json={"decision": "APPROVED"})
    assert second_attempt.status_code == 400


def test_import_requires_permission(client, db_session):
    from .conftest import auth_headers, create_user

    create_user(db_session, "viewer@example.com", "VIEWER", "Password123!")
    headers = auth_headers(client, "viewer@example.com", "Password123!")
    response = client.get("/api/v1/imports", headers=headers)
    assert response.status_code == 403
