import csv
import hashlib
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

from .mapping import SheetConfig, _norm, best_fit_config, match_config

MAX_HEADER_SCAN_ROWS = 6
MIN_HEADER_MATCHES = 2


@dataclass
class ParsedRow:
    row_number: int
    raw: dict[str, object]
    # model field name -> original formula text, only populated when this
    # sheet's SheetConfig.detect_formulas is True. See parse_workbook.
    formula_fields: dict[str, str] = field(default_factory=dict)


@dataclass
class ParsedSheet:
    name: str
    config: SheetConfig | None
    headers: list[str]
    header_row_index: int
    rows: list[ParsedRow] = field(default_factory=list)


def checksum_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _json_safe_cell(value: object) -> object:
    """A `.xlsx` date-formatted cell comes through openpyxl as a real `datetime`/
    `date` object — correct for a template with real dates, but `ImportRow.raw_data`
    is a JSON column and Python's `datetime` isn't JSON-serializable. Converted to
    its ISO string here, at the one place every raw cell value passes through,
    rather than special-cased per sheet (see validators.map_row for the matching
    fix on the *mapped* side)."""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _find_header_row(all_rows: list[tuple], config: SheetConfig | None) -> tuple[int, list[str]]:
    """Return (0-based row index, header labels) for the most likely header row.

    Templates in this workspace put a title/description row first, so the
    real header is usually row index 1 (second row); we still scan a few
    rows and score against the known column map to stay robust.
    """
    best_index = 0
    best_headers: list[str] = []
    best_score = -1
    expected = set(config.column_map.keys()) if config else set()

    for idx in range(min(MAX_HEADER_SCAN_ROWS, len(all_rows))):
        row = all_rows[idx]
        labels = [str(c).strip() if c is not None else "" for c in row]
        normalized = {_norm(label) for label in labels if label}
        score = len(normalized & expected) if expected else sum(1 for label in labels if label)
        if score > best_score:
            best_score = score
            best_index = idx
            best_headers = labels

    return best_index, best_headers


def _formula_fields_by_row(path: Path, sheet_name: str, header_index: int, headers: list[str],
                            config: SheetConfig) -> dict[int, dict[str, str]]:
    """Second pass, formulas intact (data_only=False), only for sheets that opt in
    via SheetConfig.detect_formulas — e.g. an Incident History workbook where a
    stray cross-row SUM() left in a per-incident cell must never be silently
    trusted as that row's real value (see import_service.py's incident handling).
    Returns {row_number: {model_field_name: formula_text}}."""
    wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        ws = wb[sheet_name]
        by_row: dict[int, dict[str, str]] = {}
        for row_offset, row in enumerate(ws.iter_rows(min_row=header_index + 2)):
            row_number = header_index + row_offset + 2
            for col_index, header in enumerate(headers):
                if not header or col_index >= len(row):
                    continue
                cell = row[col_index]
                if cell.data_type == "f":
                    field_name = config.column_map.get(_norm(header))
                    if field_name:
                        by_row.setdefault(row_number, {})[field_name] = str(cell.value)
        return by_row
    finally:
        wb.close()


def parse_workbook(path: Path) -> list[ParsedSheet]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parsed_sheets: list[ParsedSheet] = []

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        all_rows = [row for row in worksheet.iter_rows(values_only=True)]
        if not all_rows:
            parsed_sheets.append(ParsedSheet(name=sheet_name, config=None, headers=[], header_row_index=0))
            continue

        config = match_config(sheet_name)
        header_index, headers = _find_header_row(all_rows, config)
        if config is None:
            # xlsx sheets are normally matched purely by tab name (unlike CSV,
            # which has no tab to match and always falls back to header-overlap
            # guessing) — but a workbook whose own sheet is generically named
            # (e.g. literally "Sheet", openpyxl's own default) has no usable tab
            # name at all. Falling back to the same header-overlap guess CSV
            # already uses removes the need to rename the tab before every
            # upload, and MIN_BEST_FIT_MATCHES still guards against misrouting a
            # sheet that only coincidentally shares a column or two.
            config = best_fit_config(headers)

        formula_by_row: dict[int, dict[str, str]] = {}
        if config and config.detect_formulas:
            formula_by_row = _formula_fields_by_row(path, sheet_name, header_index, headers, config)

        rows: list[ParsedRow] = []
        for offset, data_row in enumerate(all_rows[header_index + 1 :]):
            if all(cell is None for cell in data_row):
                continue
            raw = {}
            for col_index, header in enumerate(headers):
                if not header:
                    continue
                value = data_row[col_index] if col_index < len(data_row) else None
                if isinstance(value, str):
                    value = value.strip() or None
                raw[header] = _json_safe_cell(value)
            row_number = header_index + offset + 2
            rows.append(ParsedRow(row_number=row_number, raw=raw, formula_fields=formula_by_row.get(row_number, {})))

        parsed_sheets.append(
            ParsedSheet(name=sheet_name, config=config, headers=headers, header_row_index=header_index, rows=rows)
        )

    workbook.close()
    return parsed_sheets


def parse_csv(path: Path, *, display_name: str) -> list[ParsedSheet]:
    """A CSV file has exactly one implicit "sheet" and no sheet-tab name to match
    against `match_config()` the way `.xlsx` sheet tabs are matched — `display_name`
    (the uploaded filename minus extension, e.g. "vendor_master.csv" -> "vendor_master")
    is tried first through the identical `match_config()` path every `.xlsx` sheet uses,
    matching the existing template-naming convention. If that fails (an unrecognized or
    generically-named file), `best_fit_config()` guesses from the header row's column
    overlap instead — a fallback, not the primary detection path, so a well-named file
    never depends on it."""
    with path.open(newline="", encoding="utf-8-sig") as f:
        all_rows = [tuple(row) for row in csv.reader(f)]

    if not all_rows:
        return [ParsedSheet(name=display_name, config=None, headers=[], header_row_index=0)]

    config = match_config(display_name)
    header_index, headers = _find_header_row(all_rows, config)
    if config is None:
        config = best_fit_config(headers)

    rows: list[ParsedRow] = []
    for offset, data_row in enumerate(all_rows[header_index + 1 :]):
        if all(not cell for cell in data_row):
            continue
        raw: dict[str, object] = {}
        for col_index, header in enumerate(headers):
            if not header:
                continue
            value = data_row[col_index] if col_index < len(data_row) else None
            if isinstance(value, str):
                value = value.strip() or None
            raw[header] = _json_safe_cell(value)
        rows.append(ParsedRow(row_number=header_index + offset + 2, raw=raw))

    return [ParsedSheet(name=display_name, config=config, headers=headers, header_row_index=header_index, rows=rows)]
